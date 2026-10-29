"""
Farming.py - Priorizacion de sellers para el equipo de Farming.

Farming potencia a sellers que ya venden en Paris. Cada account manager (AM)
ve una categoria (division). Esta hoja le entrega a cada AM su lista de
sellers con potencial, ordenada por venta, con el perfil de potencial, el
motivo (fortaleza + brecha, con los datos del seller) y la palanca sugerida.

Como se arma
------------
1. Una fila por seller. Un seller mixto tiene score en duros y en blandos; se
   usa el de la dureza de su categoria principal. Si no tiene score en esa
   dureza, se usa su fila con mayor venta.

2. Filtros basicos (deben cumplirse todos):
     estado_seller         in {Activo, Reactivado, Dormant}   ya vende o se apago hace poco
     categoria_reputacion  in {Medio, Bueno, Destacado, Plan Growth}   no Bajo
     sku_creados           >= 5                                catalogo minimo para trabajar
   No se filtra por segmento: entran tambien los sellers Alto.

3. Perfil de potencial (percentiles dentro de la dureza). Un seller entra si
   cumple uno de los dos (no se pueden cumplir ambos a la vez):

   Oferta  = lo que el seller pone a la venta: disponibilidad (SKU con stock, % del
             catalogo con stock) y volumen (Volumen SKU, cantidad de productos con stock).
   Demanda = lo que los clientes le compran: venta y conversion (CVR).

   Falta oferta:   FORTALEZA en demanda   Venta >= P60  o  CVR >= P60
                   BRECHA en oferta       SKU con stock < P60  o  Volumen SKU < P60
       -> sus productos se venden / convierten, pero tiene poco para vender.
          Palanca: que cargue stock y amplie catalogo (revisar lo que tiene en
          otros marketplaces), descuento en comision.

   Falta demanda:  FORTALEZA en oferta    SKU con stock >= P60  y  Volumen SKU >= P60
                   BRECHA en demanda      Venta < P40  o  CVR < P40
       -> tiene catalogo y stock, pero vende o convierte poco.
          Palanca: precio competitivo, descuento con tarjeta, SP.

4. Prioridad dentro de cada categoria: venta de mayor a menor.
   prioridad = 1 es el primer seller que deberia contactar el AM de esa categoria.
   Si se define POOL_POR_CATEGORIA, solo quedan los N primeros de cada categoria.
"""

import math

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import DUROS, BLANDOS, HORIZONTE_DIARIO
from Perfil import (CORTE_ALTO, CORTE_MEDIO_ALTO, CORTE_MEDIO, COLOR_SEGMENTO,
                    UMBRAL_PERCENTIL, UMBRAL_PERCENTIL_SKU,
                    _pct, _pesos, _num, _texto)


# ---- filtros basicos (editar aca si cambian) ----
ESTADOS_FARMING = ["Activo", "Reactivado", "Dormant"]
REPUTACION_FARMING = ["Medio", "Bueno", "Destacado", "Plan Growth"]
SKU_MIN_FARMING = 5

# ---- umbrales de los perfiles de potencial ----
UMBRAL_FORTALEZA = 60                        # variable fuerte si su percentil es >= P60
UMBRAL_BRECHA_OFERTA = UMBRAL_PERCENTIL_SKU  # stock / volumen debil si < P60 (mismo umbral que las banderas)
UMBRAL_BRECHA_DEMANDA = UMBRAL_PERCENTIL     # venta / CVR debil si < P40 (mismo umbral que las banderas)

PERFIL_OFERTA = "Falta oferta"
PERFIL_DEMANDA = "Falta demanda"
PALANCA_PERFIL = {
    PERFIL_OFERTA: "Cargar stock y ampliar catalogo (revisar otros marketplaces) + dscto en comision",
    PERFIL_DEMANDA: "Precio competitivo + dscto con tarjeta + SP",
}

POOL_POR_CATEGORIA = None   # ej: 30 -> solo los 30 primeros de cada categoria. None = todos

NOMBRE_HOJA = "Priorización de farming"
NOMBRE_HOJA_RESUMEN = "Resumen farming"

COLS_FARMING = ["categoria", "prioridad", "seller_id", "name", "admin_email",
                "perfil_farming", "motivo_farming", "palanca_farming",
                "estado_seller", "categoria_reputacion", "pct_skus_categoria", "division_db",
                "score", "segmento", "pts_para_subir_segmento",
                "flags",
                "venta", "venta_percentil",
                "tendencia_venta", "variacion_venta_30d", "venta_ult_30d", "venta_30d_anterior",
                "cvr_seller", "percentil_cvr",
                "sku_creados", "sku_stock", "pct_con_stock",
                "percentil_stock", "percentil_volumen_stock",
                "dias_vendidos", "tramos_con_venta", "percentil_recurrencia",
                "invirtio", "total_invertido"]


def _dureza_de(division) -> str:
    if division in DUROS:
        return "duros"
    if division in BLANDOS:
        return "blandos"
    return "otros"


def _pts_para_subir(fila):
    """Puntos de score que le faltan para pasar al segmento siguiente (vacio si es Alto)."""
    siguiente = {"Medio - Alto": CORTE_ALTO, "Medio": CORTE_MEDIO_ALTO, "Bajo": CORTE_MEDIO}
    corte = siguiente.get(fila["segmento"])
    if corte is None or pd.isna(fila["score"]):
        return None
    return math.ceil((corte - fila["score"]) * 10) / 10   # hacia arriba: 67.95 -> le falta 0.1


def una_fila_por_seller(df: pd.DataFrame) -> pd.DataFrame:
    """Deja una fila por seller: la de la dureza de su categoria principal."""
    df = df.copy()
    df["categoria"] = df["categoria"].fillna("Sin categoria")
    es_su_dureza = df["division_db"] == df["categoria"].apply(_dureza_de)
    df["_orden"] = es_su_dureza.astype(int)
    df = df.sort_values(["seller_id", "_orden", "venta"], ascending=[True, False, False])
    return df.drop_duplicates("seller_id").drop(columns="_orden")


def asignar_perfil(df: pd.DataFrame) -> pd.Series:
    """Perfil de potencial de cada fila: 'Falta oferta', 'Falta demanda' o vacio."""
    venta_fuerte = df["venta_percentil"] >= UMBRAL_FORTALEZA
    cvr_fuerte = df["percentil_cvr"] >= UMBRAL_FORTALEZA
    stock_fuerte = df["percentil_stock"] >= UMBRAL_FORTALEZA
    volumen_fuerte = df["percentil_volumen_stock"] >= UMBRAL_FORTALEZA

    oferta_debil = ((df["percentil_stock"] < UMBRAL_BRECHA_OFERTA)
                    | (df["percentil_volumen_stock"] < UMBRAL_BRECHA_OFERTA))
    demanda_debil = ((df["venta_percentil"] < UMBRAL_BRECHA_DEMANDA)
                     | (df["percentil_cvr"] < UMBRAL_BRECHA_DEMANDA))

    falta_oferta = (venta_fuerte | cvr_fuerte) & oferta_debil
    falta_demanda = stock_fuerte & volumen_fuerte & demanda_debil

    perfil = pd.Series("", index=df.index)
    perfil[falta_oferta] = PERFIL_OFERTA
    perfil[falta_demanda] = PERFIL_DEMANDA
    return perfil


def _pct_redondo(x, decimales: int = 0) -> str:
    """'58.12%' -> '58%' | '0.44%' con 2 decimales -> '0,44%'. Si no se puede leer, lo deja como viene."""
    try:
        valor = float(str(x).replace('%', '').replace(',', '.'))
    except ValueError:
        return _texto(x)
    return f"{valor:.{decimales}f}%".replace(".", ",")


def _motivo_farming(fila) -> str:
    """Frase para el AM: que tiene el seller (fortaleza) y que le falta (brecha).
    Cada dato va con su nombre para no confundir las dos variables de stock:
      disponibilidad = SKU con stock (% del catalogo con stock)
      volumen        = Volumen SKU (cantidad de productos con stock)
    Ej: 'Tiene oferta: buena disponibilidad (58% de su catalogo tiene stock, P60) y buen
         volumen (247 productos con stock, P93). Le falta demanda: baja conversion
         (0,44% de las visitas termina en compra, P39). Percentiles frente a sellers blandos.'"""
    dureza = fila["division_db"]
    venta = f"({_pesos(fila['venta'])} en {HORIZONTE_DIARIO} días, {_pct(fila['venta_percentil'])})"
    cvr = f"({_pct_redondo(fila['cvr_seller'], 2)} de las visitas termina en compra, {_pct(fila['percentil_cvr'])})"
    disponibilidad = (f"({_pct_redondo(fila['pct_con_stock'])} de su catálogo tiene stock, "
                      f"{_pct(fila['percentil_stock'])})")
    volumen = f"({_num(fila['sku_stock'])} productos con stock, {_pct(fila['percentil_volumen_stock'])})"

    if fila["perfil_farming"] == PERFIL_OFERTA:
        fuertes = [t for t, ok in [(f"buena venta {venta}", fila["venta_percentil"] >= UMBRAL_FORTALEZA),
                                   (f"buena conversión {cvr}", fila["percentil_cvr"] >= UMBRAL_FORTALEZA)] if ok]
        brechas = [t for t, ok in [(f"baja disponibilidad {disponibilidad}",
                                    fila["percentil_stock"] < UMBRAL_BRECHA_OFERTA),
                                   (f"poco volumen {volumen}",
                                    fila["percentil_volumen_stock"] < UMBRAL_BRECHA_OFERTA)] if ok]
        texto = f"Tiene demanda: {' y '.join(fuertes)}. Le falta oferta: {' y '.join(brechas)}."
    elif fila["perfil_farming"] == PERFIL_DEMANDA:
        brechas = [t for t, ok in [(f"baja venta {venta}", fila["venta_percentil"] < UMBRAL_BRECHA_DEMANDA),
                                   (f"baja conversión {cvr}", fila["percentil_cvr"] < UMBRAL_BRECHA_DEMANDA)] if ok]
        texto = (f"Tiene oferta: buena disponibilidad {disponibilidad} y buen volumen {volumen}. "
                 f"Le falta demanda: {' y '.join(brechas)}.")
    else:
        return ""
    return f"{texto} Percentiles frente a sellers {dureza}."


def priorizar_farming(df: pd.DataFrame) -> pd.DataFrame:
    """Sellers con potencial para Farming, con perfil y prioridad dentro de su categoria."""
    base = una_fila_por_seller(df)
    elegible = (
        base["estado_seller"].isin(ESTADOS_FARMING)
        & base["categoria_reputacion"].isin(REPUTACION_FARMING)
        & (base["sku_creados"] >= SKU_MIN_FARMING)
    )
    base = base[elegible].copy()
    base["perfil_farming"] = asignar_perfil(base)
    farm = base[base["perfil_farming"] != ""].copy()

    farm["motivo_farming"] = farm.apply(_motivo_farming, axis=1)
    farm["palanca_farming"] = farm["perfil_farming"].map(PALANCA_PERFIL)
    farm["pts_para_subir_segmento"] = farm.apply(_pts_para_subir, axis=1)

    farm = farm.sort_values(["categoria", "venta"], ascending=[True, False])
    farm["prioridad"] = farm.groupby("categoria").cumcount() + 1
    if POOL_POR_CATEGORIA:
        farm = farm[farm["prioridad"] <= POOL_POR_CATEGORIA]
    return farm


def _estilo_encabezado(ws, cols):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="D9D9D9")
    for j, col in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def _escribir_detalle(wb, tabla: pd.DataFrame):
    if NOMBRE_HOJA in wb.sheetnames:
        del wb[NOMBRE_HOJA]
    ws = wb.create_sheet(NOMBRE_HOJA)

    cols = list(tabla.columns)
    _estilo_encabezado(ws, cols)

    cell_font = Font(name="Arial", size=10)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    color_perfil = {PERFIL_OFERTA: "DDEBF7", PERFIL_DEMANDA: "FCE4D6"}   # azul claro / naranjo claro

    formatos = {"venta": "$#,##0", "total_invertido": "$#,##0", "score": "0",
                "venta_ult_30d": "$#,##0", "venta_30d_anterior": "$#,##0", "variacion_venta_30d": "0%",
                "pts_para_subir_segmento": "0.0", "pct_skus_categoria": "0%"}
    texto_largo = {"motivo_farming", "palanca_farming", "flags", "name", "admin_email"}

    for i, (_, fila) in enumerate(tabla.iterrows(), start=2):
        for j, col in enumerate(cols, start=1):
            val = fila[col]
            if pd.isna(val):
                val = None
            c = ws.cell(row=i, column=j, value=val)
            c.font = cell_font
            c.alignment = left if col in texto_largo else center
            c.border = border
            if col in formatos and val is not None:
                c.number_format = formatos[col]
        if "segmento" in cols and COLOR_SEGMENTO.get(fila["segmento"]):
            ws.cell(row=i, column=cols.index("segmento") + 1).fill = PatternFill(
                "solid", fgColor=COLOR_SEGMENTO[fila["segmento"]])
        if "perfil_farming" in cols and color_perfil.get(fila["perfil_farming"]):
            ws.cell(row=i, column=cols.index("perfil_farming") + 1).fill = PatternFill(
                "solid", fgColor=color_perfil[fila["perfil_farming"]])

    anchos = {"motivo_farming": 75, "palanca_farming": 45, "flags": 34, "name": 28,
              "admin_email": 30, "categoria": 18, "perfil_farming": 16, "tendencia_venta": 24}
    for j, col in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(j)].width = anchos.get(col, min(max(len(col), 12) + 2, 26))

    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(tabla) + 1}"
    ws.freeze_panes = "D2"   # deja fijas categoria, prioridad y seller_id


def _escribir_resumen(wb, farm: pd.DataFrame):
    """Tabla categoria x perfil (cuantos sellers con potencial tiene cada AM)."""
    if NOMBRE_HOJA_RESUMEN in wb.sheetnames:
        del wb[NOMBRE_HOJA_RESUMEN]
    ws = wb.create_sheet(NOMBRE_HOJA_RESUMEN)

    tabla = pd.crosstab(farm["categoria"], farm["perfil_farming"])
    for p in (PERFIL_OFERTA, PERFIL_DEMANDA):
        if p not in tabla.columns:
            tabla[p] = 0
    tabla = tabla[[PERFIL_OFERTA, PERFIL_DEMANDA]]
    tabla["Total"] = tabla.sum(axis=1)
    venta = farm.groupby("categoria")["venta"].sum()
    tabla["Venta de estos sellers"] = venta.reindex(tabla.index)
    tabla = tabla.sort_values("Total", ascending=False)
    total = tabla.sum(numeric_only=True)
    total.name = "Total"
    tabla = pd.concat([tabla, total.to_frame().T])

    cols = ["Categoría", PERFIL_OFERTA, PERFIL_DEMANDA, "Total", "Venta de estos sellers"]
    _estilo_encabezado(ws, cols)
    cell_font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    for i, (cat, fila) in enumerate(tabla.iterrows(), start=2):
        valores = [cat, fila[PERFIL_OFERTA], fila[PERFIL_DEMANDA], fila["Total"], fila["Venta de estos sellers"]]
        for j, val in enumerate(valores, start=1):
            c = ws.cell(row=i, column=j, value=val.item() if hasattr(val, "item") else val)
            c.font = bold if cat == "Total" else cell_font
            c.alignment = Alignment(horizontal="left" if j == 1 else "center")
            if j == 5:
                c.number_format = "$#,##0"
    for j, ancho in enumerate([22, 18, 18, 10, 24], start=1):
        ws.column_dimensions[get_column_letter(j)].width = ancho


def escribir_hoja_farming(df: pd.DataFrame, path: str = "score_growth.xlsx") -> pd.DataFrame:
    farm = priorizar_farming(df)
    tabla = farm[[c for c in COLS_FARMING if c in farm.columns]]

    wb = load_workbook(path)
    _escribir_detalle(wb, tabla)
    _escribir_resumen(wb, farm)
    wb.save(path)

    print(f"Hojas '{NOMBRE_HOJA}' y '{NOMBRE_HOJA_RESUMEN}' agregadas: "
          f"{len(farm)} sellers con potencial "
          f"({(farm['perfil_farming'] == PERFIL_OFERTA).sum()} {PERFIL_OFERTA}, "
          f"{(farm['perfil_farming'] == PERFIL_DEMANDA).sum()} {PERFIL_DEMANDA})")
    resumen_farming(farm)
    return farm


def resumen_farming(farm: pd.DataFrame):
    """Conteos en consola: cuantos sellers tiene cada AM, por perfil y segmento."""
    with pd.option_context("display.max_columns", None, "display.width", 250):
        print("\n=== FARMING: SELLERS CON POTENCIAL POR CATEGORIA Y PERFIL ===")
        print(pd.crosstab(farm["categoria"], farm["perfil_farming"], margins=True, margins_name="Total"))
        print("\n=== FARMING: PERFIL POR SEGMENTO ===")
        print(pd.crosstab(farm["perfil_farming"], farm["segmento"], margins=True, margins_name="Total"))
        print("\n=== FARMING: PERFIL POR ESTADO ===")
        print(pd.crosstab(farm["perfil_farming"], farm["estado_seller"], margins=True, margins_name="Total"))


if __name__ == "__main__":
    # Requiere que score_growth.xlsx tenga la hoja "Score Growth" con categoria, y
    # recalcula el perfilado para tener segmento y banderas.
    from Perfil import perfilar
    RUTA = "score_growth.xlsx"
    df = perfilar(pd.read_excel(RUTA, sheet_name="Score Growth"))
    escribir_hoja_farming(df, RUTA)
    print("\nListo.")
