"""
Perfil.py - Perfilado de sellers para Score Growth.

Capas:
  1. Segmento por score (cortes fijos, Alto >= 68)
  2. Banderas de brecha por variable (umbral en percentil)
  3. Perfil de banderas (texto) + conteo, para agrupar despues
  4. Palanca comercial (traduce flags -> como accionar comercialmente)
     + variable debil prioritaria y motivo con el dato que la explica

"""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import HORIZONTE_DIARIO, TRAMOS_TOTALES


# ======================================================================
# CAPA 1 - SEGMENTACION POR SCORE  (cortes fijos, continuos, sin huecos)
# ======================================================================
# Recalibrados el 27-10-2026 al pasar la frecuencia a percentil de dias.
# Los cortes anteriores (70 / 53 / 37) dejaban ~11% Alto, ~40% Medio-Alto,
# ~28% Medio y ~21% Bajo con la version mensual. Se buscaron los cortes que
# mantienen esas mismas proporciones con el score nuevo (data del 27-10-2026).
# Si cambian los pesos o las variables, hay que recalibrar.
CORTE_ALTO = 68           # >= 68
CORTE_MEDIO_ALTO = 46     # 46 a 68
CORTE_MEDIO = 32          # 32 a 46   (< 32 = Bajo)


def clasificar_segmento(score: float) -> str:
    if pd.isna(score):
        return "Sin score"
    if score >= CORTE_ALTO:
        return "Alto"
    if score >= CORTE_MEDIO_ALTO:
        return "Medio - Alto"
    if score >= CORTE_MEDIO:
        return "Medio"
    return "Bajo"


# ======================================================================
# CAPA 2 - BANDERAS DE BRECHA
# ======================================================================
# Variable debil si:
#   Venta              venta_percentil          < 40  (P40)
#   SKU con stock      percentil_stock          < 60  (P60)
#   Volumen SKU        percentil_volumen_stock  < 60  (P60)
#   Frecuencia venta   tramos_con_venta         < 2   (vendio en menos de 2 de los 3 tramos del horizonte)
#   CVR                percentil_cvr            < 40  (P40)
#   Inversion          invirtio == "no"
UMBRAL_PERCENTIL = 40
UMBRAL_PERCENTIL_SKU = 60
UMBRAL_TRAMOS = 2   # minimo de tramos (de 3) con venta para NO ser intermitente, sirve para cualquier horizonte


FLAGS = ["flag_venta_baja", "flag_stock_bajo", "flag_catalogo_bajo",
         "flag_cvr_bajo", "flag_intermitente", "flag_sin_inversion"]

ETIQUETAS = {
    "flag_venta_baja": "Venta",
    "flag_stock_bajo": "Stock",
    "flag_catalogo_bajo": "Volumen SKU",
    "flag_cvr_bajo": "CVR",
    "flag_intermitente": "Intermitente",
    "flag_sin_inversion": "SinInversion",
}


# ======================================================================
# CAPA 4 - PALANCA COMERCIAL  (traduce flags -> como accionar)
# ======================================================================
# Orden de flags por prioridad comercial (imagen "palancas").
# Empate en P1 se resuelve por este orden: Stock -> Volumen -> Frecuencia.
PRIORIDAD_FLAG = {
    "flag_stock_bajo":    1,   # P1
    "flag_catalogo_bajo": 2,   # P1 (desempate)
    "flag_intermitente":  3,   # P1 (desempate)
    "flag_venta_baja":    4,   # P2
    "flag_cvr_bajo":      5,   # P3
    "flag_sin_inversion": 6,   # P3
}

# Accionable (motivo de contacto / KPI a impactar) por flag.
# Esto es lo que da valor al mensaje: el "por que" le escribes al seller.
ACCIONABLE_FLAG = {
    "flag_stock_bajo":    "Cargar stock (60% productos con stock)",
    "flag_catalogo_bajo": "Arreglar catalogo",
    "flag_intermitente":  "Aumentar frecuencia de venta",
    "flag_venta_baja":    "Aumentar ventas (precio competitivo)",
    "flag_cvr_bajo":      "Mejorar conversion",
    "flag_sin_inversion": "Activar inversion en SP",
}

# Palanca comercial (canal) por flag. Lista de componentes para poder deduplicar.
PALANCA_FLAG = {
    "flag_stock_bajo":    ["Comunicacion masiva"],
    "flag_catalogo_bajo": ["Comunicacion masiva"],
    "flag_intermitente":  ["Dscto comercial", "Comunicacion masiva"],
    "flag_venta_baja":    ["Dscto comercial", "Comunicacion masiva"],
    "flag_cvr_bajo":      ["SP", "Comunicacion masiva", "Dscto comercial"],
    "flag_sin_inversion": ["SP", "Comunicacion masiva"],
}

# Orden en que se muestran los componentes de palanca (canal principal primero).
ORDEN_PALANCA = ["Comunicacion masiva", "Dscto comercial", "SP"]


def _flag_top(fila):
    """Flag activo de mayor prioridad, o None si no hay brechas."""
    activos = [f for f in PRIORIDAD_FLAG if fila[f]]
    if not activos:
        return None
    return min(activos, key=lambda f: PRIORIDAD_FLAG[f])


def _accionable(fila) -> str:
    """Motivo de contacto del flag prioritario (que KPI mover)."""
    flag_top = _flag_top(fila)
    if flag_top is None:
        return "Sin accion"
    return ACCIONABLE_FLAG[flag_top]


def _palanca(fila) -> str:
    """Canal comercial del flag prioritario (como contactar)."""
    flag_top = _flag_top(fila)
    if flag_top is None:
        return "Sin accion"
    return " + ".join(PALANCA_FLAG[flag_top])


# Nombre de la variable debil, como se le explica al account manager.
NOMBRE_VARIABLE = {
    "flag_stock_bajo":    "SKU con stock",
    "flag_catalogo_bajo": "Volumen SKU",
    "flag_intermitente":  "Frecuencia de venta",
    "flag_venta_baja":    "Venta",
    "flag_cvr_bajo":      "CVR",
    "flag_sin_inversion": "Inversion SP",
}


def _variable_debil(fila) -> str:
    """Variable debil prioritaria (la que hay que atacar primero)."""
    flag_top = _flag_top(fila)
    if flag_top is None:
        return "Sin variables debiles"
    return NOMBRE_VARIABLE[flag_top]


def _pesos(x) -> str:
    return "s/d" if pd.isna(x) else "$" + f"{x:,.0f}".replace(",", ".")


def _pct(x) -> str:
    return "s/d" if pd.isna(x) else f"P{x:.0f}"


def _num(x) -> str:
    return "s/d" if pd.isna(x) else f"{x:.0f}"


def _texto(x) -> str:
    return "s/d" if pd.isna(x) else str(x)


def _motivo(fila) -> str:
    """Explica con el dato del seller por que se ataca la variable prioritaria.
    Ej: 'Solo 35% de sus SKUs tiene stock (P22 en duros; umbral P60)'.
    """
    flag_top = _flag_top(fila)
    if flag_top is None:
        return "Sin variables debiles"
    dureza = fila.get("division_db", "")

    if flag_top == "flag_stock_bajo":
        return (f"Solo {_texto(fila['pct_con_stock'])} de sus SKUs tiene stock "
                f"({_pct(fila['percentil_stock'])} en {dureza}; umbral P{UMBRAL_PERCENTIL_SKU})")
    if flag_top == "flag_catalogo_bajo":
        return (f"Tiene {_num(fila['sku_stock'])} SKUs con stock "
                f"({_pct(fila['percentil_volumen_stock'])} en {dureza}; umbral P{UMBRAL_PERCENTIL_SKU})")
    if flag_top == "flag_intermitente":
        return (f"Vendio en {_num(fila['tramos_con_venta'])} de {TRAMOS_TOTALES} tramos del horizonte "
                f"({_num(fila['dias_vendidos'])} dias con venta en {HORIZONTE_DIARIO} dias)")
    if flag_top == "flag_venta_baja":
        return (f"Vendio {_pesos(fila['venta'])} en {HORIZONTE_DIARIO} dias "
                f"({_pct(fila['venta_percentil'])} en {dureza}; umbral P{UMBRAL_PERCENTIL})")
    if flag_top == "flag_cvr_bajo":
        return (f"Convierte {_texto(fila['cvr_seller'])} de sus visitas "
                f"({_pct(fila['percentil_cvr'])} en {dureza}; umbral P{UMBRAL_PERCENTIL})")
    return f"No invirtio en Sponsored Products en los ultimos {HORIZONTE_DIARIO} dias"


# ======================================================================
# PERFILADO
# ======================================================================
def perfilar(df: pd.DataFrame) -> pd.DataFrame:
    """Aplica las 4 capas y devuelve el df enriquecido."""
    df = df.copy()

    for col in ["score", "venta_percentil", "percentil_stock",
                "percentil_volumen_stock", "percentil_cvr", "recurrencia",
                "percentil_recurrencia", "tramos_con_venta"]:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")

    df["segmento"] = df["score"].apply(clasificar_segmento)

    df["flag_venta_baja"]    = df["venta_percentil"] < UMBRAL_PERCENTIL
    df["flag_stock_bajo"]    = df["percentil_stock"] < UMBRAL_PERCENTIL_SKU
    df["flag_catalogo_bajo"] = df["percentil_volumen_stock"] < UMBRAL_PERCENTIL_SKU
    df["flag_cvr_bajo"]      = df["percentil_cvr"] < UMBRAL_PERCENTIL
    df["flag_intermitente"]  = df["tramos_con_venta"] < UMBRAL_TRAMOS
    df["flag_sin_inversion"] = df["invirtio"].astype(str).str.lower() == "no"
    for f in FLAGS:
        df[f] = df[f].fillna(False)

    df["n_banderas"] = df[FLAGS].sum(axis=1)

    def armar_perfil(fila):
        activas = [ETIQUETAS[f] for f in FLAGS if fila[f]]
        return " + ".join(activas) if activas else "Sin brechas"

    df["flags"] = df.apply(armar_perfil, axis=1)

    # --- capa 4: accionable (motivo) + palanca (canal) del flag prioritario ---
    df["accionable"] = df.apply(_accionable, axis=1)
    df["palanca"] = df.apply(_palanca, axis=1)
    df["variable_debil"] = df.apply(_variable_debil, axis=1)
    df["motivo"] = df.apply(_motivo, axis=1)

    return df


# ======================================================================
# ESCRIBIR HOJA "Resumen"
# ======================================================================
COLS_RESUMEN = ["seller_id", "name", "categoria", "estado_seller", "division_db",
                "score", "segmento",
                "flags", "variable_debil", "motivo", "accionable", "palanca", "n_banderas",
                "flag_venta_baja", "flag_stock_bajo", "flag_catalogo_bajo",
                "flag_cvr_bajo", "flag_intermitente", "flag_sin_inversion"]

COLOR_SEGMENTO = {
    "Alto":       "63BE7B",   # verde
    "Medio - Alto":        "A9D08E",   # verde claro
    "Medio": "FFEB84",   # amarillo
    "Bajo":        "F8696B",   # rojo
    "Sin score":     "D9D9D9",   # gris
}


def escribir_hoja_resumen(df: pd.DataFrame, path: str = "score_growth.xlsx"):
    resumen = df[[c for c in COLS_RESUMEN if c in df.columns]].copy()

    for f in FLAGS:
        if f in resumen.columns:
            resumen[f] = resumen[f].map({True: "SI", False: ""})

    resumen = resumen.sort_values(["division_db", "score"], ascending=[True, False])

    wb = load_workbook(path)
    if "Resumen" in wb.sheetnames:
        del wb["Resumen"]
    ws = wb.create_sheet("Resumen")

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    cell_font = Font(name="Arial", size=10)
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    cols = list(resumen.columns)

    for j, col in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center
        c.border = border

    seg_idx = cols.index("segmento") + 1 if "segmento" in cols else None

    for i, (_, fila) in enumerate(resumen.iterrows(), start=2):
        for j, col in enumerate(cols, start=1):
            val = fila[col]
            if pd.isna(val):
                val = None
            c = ws.cell(row=i, column=j, value=val)
            c.font = cell_font
            c.alignment = center
            c.border = border
            if col == "score" and val is not None:
                c.number_format = "0"

        if seg_idx:
            seg_val = fila["segmento"]
            color = COLOR_SEGMENTO.get(seg_val)
            if color:
                ws.cell(row=i, column=seg_idx).fill = PatternFill("solid", fgColor=color)

    n_filas = len(resumen) + 1

    for j, col in enumerate(cols, start=1):
        letra = get_column_letter(j)
        ancho = min(max(len(str(col)), 12) + 2, 32)
        if col == "flags":
            ancho = 30
        if col == "accionable":
            ancho = 38
        if col == "palanca":
            ancho = 32
        if col == "motivo":
            ancho = 60
        ws.column_dimensions[letra].width = ancho

    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n_filas}"
    ws.freeze_panes = "A2"

    wb.save(path)
    print(f"Hoja 'Resumen' agregada a: {path}  ({len(resumen)} filas)")


def resumen_consola(df: pd.DataFrame):
    """Conteos rapidos por segmento y por perfil, en consola."""
    print("\n=== SELLERS POR SEGMENTO Y DUREZA ===")
    print(pd.crosstab(df["segmento"], df["division_db"]))

    print("\n=== PERFILES DE BRECHA MAS COMUNES (top 15) ===")
    print(df["flags"].value_counts().head(15))

    print("\n=== ACCIONABLE PRIORITARIO MAS COMUN ===")
    print(df["accionable"].value_counts())


if __name__ == "__main__":
    RUTA = "score_growth.xlsx"
    df = pd.read_excel(RUTA)

    df = perfilar(df)
    resumen_consola(df)
    escribir_hoja_resumen(df, RUTA)
    print("\nListo.")