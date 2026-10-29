"""
Tratamientos.py - Asignacion de tratamiento comercial a sellers.

Genera dos listas filtradas a partir de la tabla larga (con score + perfilado):
  - Regalo SP   -> sellers que reciben Sponsored Products gratis
  - Personalizado BD -> sellers con tratamiento personalizado (Business Development)


Criterios
---------
Con horizonte diario, los "meses" se reemplazan por tramos: el horizonte se divide
siempre en 3 tramos iguales (con 90 dias, de 30 dias cada uno). Ver tramos_con_venta
en Recurrencia_venta.py.

Regalo SP  (basta cumplir Regla A O Regla B):

  Regla A - Dormant:
    segmento in {Medio - Alto, Medio, Bajo}
    estado_seller == Dormant
    tramos_con_venta >= 2        (vendio en 2 de 3 tramos; antes "2 de 3 meses")
    sku_creados >= 5
    categoria_reputacion in {Medio, Bueno, Destacado, Plan Growth}

  Regla B - Activo:
    segmento in {Alto, Medio - Alto}
    estado_seller == Activo
    tramos_con_venta >= 3        (vendio en los 3 tramos; antes "los 3 meses")
    percentil_cvr >= 60
    sku_creados >= 5
    categoria_reputacion in {Medio, Bueno, Destacado, Plan Growth}

Personalizado BD  (basta cumplir Regla A O Regla B):

  Regla A - Activo:
    segmento == Alto
    estado_seller == Activo
    categoria_reputacion in {Bueno, Destacado, Plan Growth}

  Regla B - Dormant:
    segmento == Medio - Alto
    estado_seller == Dormant
    categoria_reputacion in {Bueno, Plan Growth}
"""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ---- parametros de los criterios (editar aca si cambian) ----
REPUTACION_OK = ["Bueno", "Destacado", "Plan Growth"]
REPUTACION_SP = ["Medio","Bueno", "Destacado", "Plan Growth"]
REPUTACION_BD_DORMANT = ["Bueno", "Plan Growth"]
SEGMENTOS_SP_DORMANT = ["Medio - Alto", "Medio", "Bajo"]   # Regla A de Regalo SP
SKU_MIN = 5
TRAMOS_MIN_DORMANT = 2   # vendio en al menos 2 de los 3 tramos del horizonte
TRAMOS_MIN_ACTIVO = 3    # vendio en los 3 tramos del horizonte
PERCENTIL_CVR_MIN = 60

COLS_TRATAMIENTO = ["seller_id", "name", "estado_seller", "admin_email", "categoria", "categoria_reputacion",
                    "division_db", "score", "segmento",
                    "venta", "venta_percentil",
                    "dias_vendidos", "tramos_con_venta", "recurrencia", "percentil_recurrencia",
                    "cvr_seller", "percentil_cvr",
                    "sku_creados", "sku_stock", "pct_con_stock",
                    "percentil_stock", "percentil_volumen_stock",
                    "invirtio", "total_invertido"]


def filtrar_sp(df: pd.DataFrame) -> pd.DataFrame:
    """Sellers que reciben regalo de Sponsored Products (Regla A o Regla B)."""
    rep_ok = df["categoria_reputacion"].isin(REPUTACION_SP)
    sku_ok = df["sku_creados"] >= SKU_MIN

    # Regla A - Dormant
    regla_a = (
        df["segmento"].isin(SEGMENTOS_SP_DORMANT)
        & (df["estado_seller"] == "Dormant")
        & (df["tramos_con_venta"] >= TRAMOS_MIN_DORMANT)
        & sku_ok
        & rep_ok
    )

    # Regla B - Activo
    regla_b = (
        df["segmento"].isin(["Alto", "Medio - Alto"])
        & (df["estado_seller"] == "Activo")
        & (df["tramos_con_venta"] >= TRAMOS_MIN_ACTIVO)
        & (df["percentil_cvr"] >= PERCENTIL_CVR_MIN)
        & sku_ok
        & rep_ok
    )

    sp = df[regla_a | regla_b].copy()
    sp["regla_sp"] = ""
    sp.loc[regla_a[sp.index], "regla_sp"] = "A-Dormant"
    sp.loc[regla_b[sp.index], "regla_sp"] = "B-Activo"
    sp.loc[(regla_a & regla_b)[sp.index], "regla_sp"] = "A+B"
    return sp


def filtrar_bd(df: pd.DataFrame) -> pd.DataFrame:
    """Sellers con tratamiento personalizado (Business Development): Regla A o Regla B."""
    # Regla A - Activo Alto
    regla_a = (
        (df["segmento"] == "Alto")
        & (df["estado_seller"] == "Activo")
        & (df["categoria_reputacion"].isin(REPUTACION_OK))
    )
    # Regla B - Dormant Medio - Alto
    regla_b = (
        (df["segmento"] == "Medio - Alto")
        & (df["estado_seller"] == "Dormant")
        & (df["categoria_reputacion"].isin(REPUTACION_BD_DORMANT))
    )
    bd = df[regla_a | regla_b].copy()
    bd["regla_bd"] = ""
    bd.loc[regla_a[bd.index], "regla_bd"] = "A-Activo"
    bd.loc[regla_b[bd.index], "regla_bd"] = "B-Dormant"
    return bd


def _escribir_hoja(wb, nombre_hoja: str, datos: pd.DataFrame, extra_cols=None):
    cols = [c for c in COLS_TRATAMIENTO if c in datos.columns]
    if extra_cols:
        cols = cols + [c for c in extra_cols if c in datos.columns]
    tabla = datos[cols].copy()
    tabla = tabla.sort_values(["division_db", "score"], ascending=[True, False])

    if nombre_hoja in wb.sheetnames:
        del wb[nombre_hoja]
    ws = wb.create_sheet(nombre_hoja)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    cell_font = Font(name="Arial", size=10)
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    lista_cols = list(tabla.columns)

    for j, col in enumerate(lista_cols, start=1):
        c = ws.cell(row=1, column=j, value=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center
        c.border = border

    for i, (_, fila) in enumerate(tabla.iterrows(), start=2):
        for j, col in enumerate(lista_cols, start=1):
            val = fila[col]
            if pd.isna(val):
                val = None
            c = ws.cell(row=i, column=j, value=val)
            c.font = cell_font
            c.alignment = center
            c.border = border
            if col == "score" and val is not None:
                c.number_format = "0"

    n_filas = len(tabla) + 1
    for j, col in enumerate(lista_cols, start=1):
        letra = get_column_letter(j)
        ancho = min(max(len(str(col)), 12) + 2, 30)
        ws.column_dimensions[letra].width = ancho

    ws.auto_filter.ref = f"A1:{get_column_letter(len(lista_cols))}{n_filas}"
    ws.freeze_panes = "A2"
    print(f"Hoja '{nombre_hoja}' agregada: {len(tabla)} filas")


def escribir_hojas_tratamiento(df: pd.DataFrame, path: str = "score_growth.xlsx"):
    """Filtra y escribe las hojas 'Regalo SP' y 'Personalizado BD'."""
    sp = filtrar_sp(df)
    bd = filtrar_bd(df)

    wb = load_workbook(path)
    _escribir_hoja(wb, "Regalo SP", sp, extra_cols=["regla_sp"])
    _escribir_hoja(wb, "Personalizado BD", bd, extra_cols=["regla_bd"])
    wb.save(path)

    print(f"\nResumen tratamientos:")
    print(f"  Regalo SP:        {len(sp)} filas  ({sp['seller_id'].nunique()} sellers unicos)")
    print(f"  Personalizado BD: {len(bd)} filas  ({bd['seller_id'].nunique()} sellers unicos)")
    return sp, bd


if __name__ == "__main__":
    RUTA = "score_growth.xlsx"
    df = pd.read_excel(RUTA, sheet_name="Score Growth")   # <- Score Growth, no Resumen
    escribir_hojas_tratamiento(df, RUTA)
    print("\nListo.")