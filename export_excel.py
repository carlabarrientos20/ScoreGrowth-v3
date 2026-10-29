"""
export_excel.py - Exporta la tabla LARGA del Score Growth a Excel con formato
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter


COLS_PERCENTIL = ["venta_percentil", "percentil_cvr", "percentil_stock", "percentil_volumen_stock",
                  "percentil_recurrencia", "score"]
COLS_PLATA = ["venta", "total_invertido", "venta_ult_30d", "venta_30d_anterior"]
COLS_PCT = ["margen_pct", "variacion_venta_30d"]
COLS_SEMAFORO = ["dias_vendidos", "tramos_con_venta"]


def exportar(df: pd.DataFrame, path: str = "score_growth.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Score Growth"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    cell_font = Font(name="Arial", size=10)
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    cols = list(df.columns)
    if len(df) == 0:
        print("⚠️ DataFrame vacío, no se genera Excel. Alguna variable vino sin datos.")
        return

    # encabezados
    for j, col in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center
        c.border = border

    # datos
    for i, (_, fila) in enumerate(df.iterrows(), start=2):
        for j, col in enumerate(cols, start=1):
            val = fila[col]
            if pd.isna(val):
                val = None
            c = ws.cell(row=i, column=j, value=val)
            c.font = cell_font
            c.alignment = center
            c.border = border
            if col in COLS_PLATA and val is not None:
                c.number_format = '$#,##0'
            elif col in COLS_PCT and val is not None:
                c.number_format = '0.0%'
            elif col in COLS_PERCENTIL and val is not None:
                c.number_format = '0'
            elif col in COLS_SEMAFORO and val is not None:
                c.number_format = '0'

    n_filas = len(df) + 1


    for col in COLS_PERCENTIL:
        if col in cols:
            letra = get_column_letter(cols.index(col) + 1)
            rango = f"{letra}2:{letra}{n_filas}"
            regla = ColorScaleRule(
                start_type="num", start_value=0, start_color="F8696B",
                mid_type="num", mid_value=50, mid_color="FFEB84",
                end_type="num", end_value=100, end_color="63BE7B",
            )
            ws.conditional_formatting.add(rango, regla)

  
    for col in COLS_SEMAFORO:
        if col in cols:
            letra = get_column_letter(cols.index(col) + 1)
            rango = f"{letra}2:{letra}{n_filas}"
            regla = ColorScaleRule(
                start_type="min", start_color="F8696B",
                mid_type="percentile", mid_value=50, mid_color="FFEB84",
                end_type="max", end_color="63BE7B",
            )
            ws.conditional_formatting.add(rango, regla)


    for j, col in enumerate(cols, start=1):
        letra = get_column_letter(j)
        ws.column_dimensions[letra].width = min(max(len(str(col)), 14) + 2, 30)

    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n_filas}"
    ws.freeze_panes = "A2"

    wb.save(path)
    print(f"Excel guardado en: {path}")

if __name__ == "__main__":
    from main import tabla_larga
    exportar(tabla_larga(), "score_growth.xlsx")